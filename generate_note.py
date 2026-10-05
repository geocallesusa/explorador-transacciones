from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

filename = 'Nota_tecnica_breve.docx'

doc = Document()
section = doc.sections[0]
section.top_margin = 720
section.bottom_margin = 720
section.left_margin = 720
section.right_margin = 720

paragraph = doc.add_paragraph()
run = paragraph.add_run('Nota técnica breve')
run.bold = True
run.font.size = Pt(18)
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

content = [
    (
        'Decisiones de diseño',
        'Se optó por una arquitectura simple con separación clara entre la interfaz y la lógica de negocio. La UI se encarga de cargar archivos, mostrar métricas y filtrar resultados, mientras que el motor de normalización centraliza la validación, detección de fuentes, limpieza de campos y rechazo de registros inconsistentes. También se priorizó un flujo local en navegador, sin backend, para mantener la solución ligera y fácil de probar.'
    ),
    (
        'Uso de IA',
        'La IA se utilizó para acelerar la definición del alcance, generar la estructura del proyecto, proponer la lógica de normalización y sugerir criterios de validación para distintos formatos JSON. Esto permitió avanzar rápidamente en la base funcional y mantener una visión clara del comportamiento esperado.'
    ),
    (
        'Ajustes sobre sugerencias de IA',
        'Se realizaron ajustes para alinear la solución con los requisitos del proyecto: se estandarizó el idioma del proyecto al español, se renombraron los archivos de prueba, se agregó una vista explícita de transacciones rechazadas con motivo de rechazo, y se validó la compilación final con build real. Además, se corrigieron detalles de normalización y nomenclatura para evitar inconsistencias de uso y mejorar la legibilidad del producto final.'
    ),
]

for heading, text in content:
    heading_paragraph = doc.add_paragraph()
    heading_run = heading_paragraph.add_run(heading)
    heading_run.bold = True
    heading_run.font.size = Pt(12)
    doc.add_paragraph(text)

# Aseguramos una hoja breve y clara.
doc.save(filename)
print(f'Archivo creado: {filename}')
